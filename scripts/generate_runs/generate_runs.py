import copy
import os
import sys
from openpyxl.reader.excel import load_workbook
import xml.etree

parameters_filename = 'calibration parameters.xlsx'
post_calib_filename = 'post calib.xlsx'
earliest_end_year = 2030
'''
The interventions parameters can be modified by setting the following macros.

INTERVENTIONS is a global switch for the parameters: set it to false to skip all intervention substitutions.
The _MEAN_MULTIPLIER macros are simple multipliers for the variables.
The _STDEV_MODIFIER macros can have two different behaviors: 
	1) if _IS_NUMBER is True, then the modifier is considered just a number
	and is substituted in all files "as is".
	2) if _IS_PERCENTAGE is true, it will act as a multiplier of the mean:
	for example you can set it to 0.5 to have the standard deviation be
	half of the mean for that variable.
	(Note that setting them both to True or False will cause an error.)
'''
INTERVENTIONS = True
ACQ_RATE_MEAN_MULTIPLIER = 2.0
ACQ_RATE_STDEV_IS_NUMBER = False
ACQ_RATE_STDEV_IS_PERCENTAGE = True
ACQ_RATE_STDEV_MODIFIER= 0.5
REG_ACTS_MEAN_MULTIPLIER = 2.0

header = ['Assort', 'PropHRMale', 'PropHRFemale', 'epsilon', 'HRMult', 'CSWMult', 'RegActs', 'ChanceCSW', 'AqRateStdyLR', 'AqRateRegLR', 'AqRateCasLR', 'AqRateCSWLR', 'AqRateStdyHR', 'AqRateRegHR', 'AqRateCasHR', 'AqRateCSWHR']

def read_months_of_1990():
	print('Reading months of 1990...', end='')
	workbook = load_workbook(filename = post_calib_filename)
	#sheet_ranges = workbook.get_sheet_by_name(name='Sheet1')
	sheet_ranges = workbook.active
	months = {row[0].value : int(row[1].value) for row in sheet_ranges.rows[1:]}
	print('done.')
	return months

def lookup_month_of_1990(starts_with, months_of_1990):
	for name in months_of_1990:
		if name.startswith(starts_with):
			return months_of_1990[name]
	return None

def read_parameters(parameters_filename):
	print('Reading parameters...', end='')
	workbook = load_workbook(filename = parameters_filename)
	sheet_ranges = workbook.active
	#sheet_ranges = workbook.get_sheet_by_name(name='Sheet1')
	rows = [[cell.value for cell in row] for row in sheet_ranges.rows[1:]]
	print('done.')
	return rows

def read_template(template_filename):
	print('Reading template...', end='')
	tree = xml.etree.ElementTree.parse(template_filename)
	print('done.')
	return tree

def filter_rows_by_weight(rows, cutoff):
	sum_weight = 0
	parameter_sets = {}

	print('Filtering runs by weight...', end='')
	for row in sorted(rows, key=lambda r: float(r[1]), reverse=True):
		sum_weight += float(row[1])
		parameter_sets[row[0]] = [float(i) for i in row[2:]]
		if sum_weight > cutoff:
			break
	print('done.')
	return parameter_sets

def try_make_directory(directory):
	try:
		os.mkdir(directory)
	except:
		print()
		print('Error: output folder already exists - {}'.format(directory))
		return False
	return True

def replace_parameters(parameters, template_tree, month_of_1990):
	tree = copy.deepcopy(template_tree)
	root = tree.getroot()

	root.find('./monthOf1990').text = str(month_of_1990)

	if earliest_end_year > 0:
		duration = int(root.find('./duration').text)
		months_after_1990 = duration - month_of_1990
		if (1990 + months_after_1990 / 12) < earliest_end_year:
			root.find('./duration').text = str(month_of_1990 + 12 * (earliest_end_year - 1990))

	if root.find('./interventions/populationInterventions'):
		for intervention in root.find('./interventions/populationInterventions'):
			adjusted_time = month_of_1990 + 12 * (int(intervention.attrib['time']) - 1990)
			intervention.set('time', str(adjusted_time))

	if root.find('./interventions/groups'):
		for group in root.find('./interventions/groups'):
			period = group.find('./enrollment-period').text
			if '-' in period:
				start, end = map(int, period.split('-'))
				start = month_of_1990 + 12 * (start - 1990)
				end = month_of_1990 + 12 * (end - 1990)
				group.find('./enrollment-period').text = '{}-{}'.format(start, end)
			else:
				time = month_of_1990 + 12 * (int(period) - 1990)
				group.find('./enrollment-period').text = '{}-{}'.format(time)

	for rollout_file_time in root.findall('./interventions/artRolloutIntervention/rolloutTreatmentFiles/rolloutFile/time'):
		if int(rollout_file_time.text) in [-1, 0]:
			continue
		adjusted_time = month_of_1990 + 12 * (int(rollout_file_time.text) - 1990)
		rollout_file_time.text = str(adjusted_time)

	replacement_path_map = {
		'./population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Steady"]/assortativeness': 'Assort',
		'./population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Regular"]/assortativeness': 'Assort',
		'./population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Casual"]/assortativeness': 'Assort',
		'./population/entities/entity[@type="Male"]/behavior/proportionHighRiskNonCsw' : 'PropHRMale',
		'./population/entities/entity[@type="Male"]/behavior/proportionHighRiskCsw' : 'PropHRMale',
		'./population/entities/entity[@type="Male"]/behavior/partnerAcqMultWithSteadyLowRisk' : 'epsilon',
		'./population/entities/entity[@type="Male"]/behavior/highRiskAcqRateMultiplier' : 'HRMult',
		'./population/entities/entity[@type="Male"]/behavior/highRiskCswAcqRateMultiplier' : 'CSWMult',
		'./population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Regular"]/coitalEventsPerMonthLowRisk/distribution/mean' : 'RegActs',
		'./population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Regular"]/coitalEventsPerMonthHighRisk/distribution/mean' : 'RegActs',
		'./population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Steady"]/acquisitionRateLowRisk/distribution/mean' : 'AqRateStdyLR',
		'./population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Steady"]/acquisitionRateHighRisk/distribution/mean' : 'AqRateStdyHR',
		'./population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Regular"]/acquisitionRateLowRisk/distribution/mean' : 'AqRateRegLR',
		'./population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Regular"]/acquisitionRateHighRisk/distribution/mean' : 'AqRateRegHR',
		'./population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Casual"]/acquisitionRateLowRisk/distribution/mean' : 'AqRateCasLR',
		'./population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Casual"]/acquisitionRateHighRisk/distribution/mean' : 'AqRateCasHR',
		'./population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Csw"]/acquisitionRateLowRisk/distribution/mean' : 'AqRateCSWLR',
		'./population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Csw"]/acquisitionRateHighRisk/distribution/mean' : 'AqRateCSWHR',
		'./population/entities/entity[@type="Female"]/behavior/proportionHighRiskNonCsw' : 'PropHRFemale',
		'./population/initialState/chanceBeingCswFemale' : 'ChanceCSW',
		'./population/entities/entity[@type="Female"]/behavior/chanceBecomeSexWorker' : 'ChanceCSW'
	}

	# Added to support  modifying the interventions with a multiplier
	acq_rate_intervention_path_map = {
		'.//partitions/partition/interventions/partnerAcquisitionRate[@risk="low"][@type="steady"]' : './population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Steady"]/acquisitionRateLowRisk/distribution/mean',
		'.//partitions/partition/interventions/partnerAcquisitionRate[@risk="high"][@type="steady"]' : './population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Steady"]/acquisitionRateHighRisk/distribution/mean',
		'.//partitions/partition/interventions/partnerAcquisitionRate[@risk="low"][@type="regular"]' : './population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Regular"]/acquisitionRateLowRisk/distribution/mean',
		'.//partitions/partition/interventions/partnerAcquisitionRate[@risk="high"][@type="regular"]' : './population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Regular"]/acquisitionRateHighRisk/distribution/mean',
		'.//partitions/partition/interventions/partnerAcquisitionRate[@risk="low"][@type="casual"]' : './population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Casual"]/acquisitionRateLowRisk/distribution/mean',
		'.//partitions/partition/interventions/partnerAcquisitionRate[@risk="high"][@type="casual"]' : './population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Casual"]/acquisitionRateHighRisk/distribution/mean',
		'.//partitions/partition/interventions/partnerAcquisitionRate[@risk="low"][@type="csw"]' : './population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Csw"]/acquisitionRateLowRisk/distribution/mean',
		'.//partitions/partition/interventions/partnerAcquisitionRate[@risk="high"][@type="csw"]' : './population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Csw"]/acquisitionRateHighRisk/distribution/mean'
	}

	reg_acts_intervention_path_map = {
		'.//partitions/partition/interventions/coitalEventsPerMonth[@type="regular"][@risk="low"]' : './population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Regular"]/coitalEventsPerMonthLowRisk/distribution/mean',
		'.//partitions/partition/interventions/coitalEventsPerMonth[@type="regular"][@risk="high"]' : './population/entities/entity[@type="Male"]/behavior/partnershipTypes/partnership[@type="Regular"]/coitalEventsPerMonthHighRisk/distribution/mean'
	}

	for path in replacement_path_map:
		parameter = replacement_path_map[path]
		parameter_element = root.find(path)
		parameter_element.text = str(parameters[parameter])

	 # Handles intervention substitutions: set INTERVENTIONS = False at the beginning of the file to disable
	if (INTERVENTIONS): 
		assert (ACQ_RATE_STDEV_IS_PERCENTAGE ^ ACQ_RATE_STDEV_IS_NUMBER), \
		"Check intervention parameters: the standard deviation modifier needs to be set as percentage OR number."

		for path in acq_rate_intervention_path_map:
			if root.find(path):
				parameter_path = root.find(path)
				parameter_element_mean = parameter_path.find("distribution/mean")
				parameter_element_mean.text = str(float(root.find(acq_rate_intervention_path_map[path]).text) * ACQ_RATE_MEAN_MULTIPLIER)
				parameter_element_stDev = parameter_path.find("distribution/stdDev")
				if (ACQ_RATE_STDEV_IS_PERCENTAGE):
					parameter_element_stDev.text = str(float(parameter_element_mean.text) * ACQ_RATE_STDEV_MODIFIER)
				elif (ACQ_RATE_STDEV_IS_NUMBER):
					parameter_element_stDev.text = str(ACQ_RATE_STDEV_MODIFIER)

		for path in reg_acts_intervention_path_map:
			if root.find(path):
				parameter_path = root.find(path)
				parameter_element_mean = parameter_path.find("distribution/mean")
				parameter_element_mean.text = str(round(float(root.find(reg_acts_intervention_path_map[path]).text) * REG_ACTS_MEAN_MULTIPLIER))
				# We skip the standard deviation here because the distribution is a Poisson, thus \mu = \sigma^2 = \lambda
	return tree

def write_xml_files(parameter_sets, template, output_directory, months_of_1990, suffix=None):
	for run in parameter_sets:
		if suffix:
			xml_filename = os.path.join(output_directory, run  + '_' + suffix + '.xml')
		else:
			xml_filename = os.path.join(output_directory, run + '.xml')
		parameters = {k : v for k, v in zip(header, parameter_sets[run])}
		month_of_1990 = lookup_month_of_1990(run, months_of_1990)

		if not month_of_1990:
			print('Month of 1990 not found for: ' + run)
			print('Skipping...')
			continue

		new_xml = replace_parameters(parameters, template, month_of_1990)
		new_xml.write(xml_filename)

def write_xml_batches(parameter_sets, template, base_output_directory, batch_size, months_of_1990, suffix=None):
	keys = parameter_sets.keys()

	for i in range(0, len(parameter_sets), batch_size):
		subset_keys = list(keys)[i:i+batch_size]
		batch_parameter_sets = {j : parameter_sets[j] for j in subset_keys}
		batch_directory = os.path.join(base_output_directory, 'batch' + str(i // batch_size))

		if not try_make_directory(batch_directory):
			return

		write_xml_files(batch_parameter_sets, template, batch_directory, months_of_1990, suffix)

def generate_runs(template_filename, weight_cutoff, run_set_name, batch_size):
	months_of_1990 = read_months_of_1990()
	rows = read_parameters(parameters_filename)
	parameter_sets = filter_rows_by_weight(rows, weight_cutoff)
	template = read_template(template_filename)

	generate_message = 'Generating XML files for the top {:g}% of parameter sets by weight ({} files)...'
	print(generate_message.format(weight_cutoff * 100, len(parameter_sets)), end='')

	base_output_directory = os.path.join(os.path.dirname(parameters_filename), run_set_name)

	if not try_make_directory(base_output_directory):
		return

	if batch_size == 0:
		write_xml_files(parameter_sets, template, base_output_directory, months_of_1990, run_set_name)
	else:
		write_xml_batches(parameter_sets, template, base_output_directory, batch_size, months_of_1990, run_set_name)

	print('done.')
	print('XMLs can be found in the folder {}.'.format(base_output_directory))

def run(args):
	if len(args) < 3 or len(args) > 5:
		print('Usage: generate_runs.py [batch_size=0] [weight_cutoff=0.9] template_file run_set_name')
		return

	args.pop(0)

	batch_size = 0
	weight_cutoff = 0.9

	while len(args) > 2:
		arg = args.pop(0)
		if arg.startswith('batch_size='):
			batch_size = int(arg[11:])
		elif arg.startswith('weight_cutoff='):
			weight_cutoff = float(arg[15:])
		else:
			print('unknown argument {} ignored'.format(arg))

	generate_runs(args[0], weight_cutoff, args[1], batch_size)

if __name__ == '__main__':
	run(sys.argv)
